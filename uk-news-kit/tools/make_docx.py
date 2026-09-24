"""將 post.md 轉做 Word 檔（連圖），俾人手覆核。

用法：python3 tools/make_docx.py posts-uk/YYYY-MM-DD/post.md
輸出：同一個資料夾嘅 post.docx
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def main():
    src = Path(sys.argv[1])
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft JhengHei"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")

    doc.add_heading(src.parent.name + " 英國新聞稿", level=0)
    for line in src.read_text(encoding="utf-8").splitlines():
        img = re.match(r"^!\[.*?\]\((.+?)\)", line.strip())
        if img:
            path = src.parent / img.group(1)
            if path.exists():
                doc.add_picture(str(path), width=Cm(12))
            else:
                doc.add_paragraph(f"（搵唔到圖：{img.group(1)}）")
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif re.match(r"^\s*[-*] ", line):
            doc.add_paragraph(re.sub(r"^\s*[-*] ", "", line), style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)
    out = src.with_suffix(".docx")
    doc.save(out)
    print(f"OK {out}")


if __name__ == "__main__":
    main()
